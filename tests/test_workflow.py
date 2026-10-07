import io
import time
import pytest
from PIL import Image
from openpyxl import load_workbook, Workbook
from fastapi.testclient import TestClient
import app as module
from tables import from_textin, from_workbook, export, validate


def photo():
    data=io.BytesIO()
    Image.new('RGB',(320,200),'white').save(data,'PNG')
    return data.getvalue()


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(module,'DATA',tmp_path)
    return TestClient(module.app)


def cell(**extra):
    c=dict(id='0',row=1,col=1,rowspan=1,colspan=1,text='0012',score=.99,polygon=[],confirmed=False,kind='cell')
    c.update(extra)
    return c


def test_literals_and_original(tmp_path):
    image=tmp_path/'original.png'
    Image.new('RGB',(50,50),'white').save(image)
    path=tmp_path/'result.xlsx'
    cells=[cell(id=str(i),row=i+1,text=t,colspan=2 if i==0 else 1) for i,t in enumerate(['标题','0012','=1+1','-0.25','@SUM(A1)',''])]
    export(cells,image,path)
    wb=load_workbook(path,data_only=False)
    assert wb.sheetnames==['识别表格','原图']
    assert len(wb.worksheets[1]._images)==1
    assert str(next(iter(wb.active.merged_cells.ranges)))=='A1:B1'
    for i,c in enumerate(cells):
        raw=wb.active.cell(i+1,1)
        assert (raw.value or '')==(-0.25 if c['text']=='-0.25' else c['text'])
        assert raw.data_type!='f'
        assert raw.comment is None
        assert raw.alignment.horizontal==raw.alignment.vertical=='center'


def test_overlap_rejected():
    with pytest.raises(ValueError,match='重叠'):
        validate([cell(colspan=2),cell(id='1',col=2)])


def test_numeric_export_preserves_identifiers_and_expressions(tmp_path):
    image=tmp_path/'original.png'
    Image.new('RGB',(50,50),'white').save(image)
    texts=['8','15','0','-3','12.50','0012','012.5','=1+1','1+2','1 4',
           '1234567890123456','客户16801','+2','0.0000000000000000001']
    target=tmp_path/'numbers.xlsx'
    export([cell(id=str(i),row=i+1,text=t) for i,t in enumerate(texts)],image,target)
    ws=load_workbook(target).active
    for row,value in {1:8,2:15,3:0,4:-3,5:12.5,13:2,14:1e-19}.items():
        assert ws.cell(row,1).value==value and ws.cell(row,1).data_type=='n'
    assert ws['A5'].number_format=='0.00'
    for row in range(6,13):
        assert ws.cell(row,1).value==texts[row-1] and ws.cell(row,1).data_type=='s'
    assert sum(ws.cell(row,1).value for row in (1,2))==23
    assert all(ws.cell(row,1).data_type!='f' for row in range(1,len(texts)+1))


def test_history_delete_restore_keeps_original(client,tmp_path):
    j=client.post('/api/jobs',files={'file':('history.png',photo(),'image/png')}).json()
    jid=j['id']
    original=(tmp_path/jid/'source.upload').read_bytes()
    response=client.delete('/api/jobs/'+jid)
    assert response.status_code==200 and response.json()['recoverable']
    assert not client.get('/api/jobs').json()
    assert client.get('/api/jobs/'+jid).status_code==404
    assert client.post('/api/trash/'+jid+'/restore').status_code==200
    assert (tmp_path/jid/'source.upload').read_bytes()==original
    assert client.get('/api/jobs').json()[0]['id']==jid


def test_history_delete_rejects_active_and_bad_ids(client,tmp_path):
    j=client.post('/api/jobs',files={'file':('history.png',photo(),'image/png')}).json()
    path=tmp_path/j['id'];j['status']='recognizing';module.write_job(path,j)
    assert client.delete('/api/jobs/'+j['id']).status_code==409
    assert path.exists()
    assert client.post('/api/trash/not-a-job/restore').status_code==404


def test_textin_current_structure_and_scores():
    payload={'code':200,'result':{'tables':[
        {'position':[0,0,200,0,200,20,0,20],'lines':[{'text':'当前标题','score':.98,'position':[0,0,200,0,200,20,0,20]}]},
        {'position':[0,30,200,30,200,90,0,90],'table_cells':[
            {'start_row':1,'end_row':1,'start_col':1,'end_col':2,'text':'新表头','position':[0,30,200,30,200,50,0,50],'lines':[{'score':.96}]},
            {'start_row':2,'end_row':2,'start_col':1,'end_col':1,'text':'0012','position':[0,50,100,50,100,90,0,90],'lines':[{'score':.9}]},
            {'start_row':2,'end_row':2,'start_col':2,'end_col':2,'text':'','position':[100,50,200,50,200,90,100,90],'lines':[]}]}
    ]}}
    cells,warnings=from_textin(payload)
    assert [c['text'] for c in cells]==['当前标题','新表头','0012','']
    assert [c['row'] for c in cells]==[1,2,3,3]
    assert cells[1]['colspan']==2
    assert cells[2]['score']==.9 and cells[3]['score'] is None
    assert warnings
    with pytest.raises(ValueError,match='TextIn 错误'):
        from_textin({'code':40003,'message':'余额不足'})


def test_cloud_formula_cache_and_number_format(tmp_path):
    wb=Workbook()
    wb.active['A1']=12
    wb.active['A1'].number_format='0000'
    wb.active['A2']='=1+1'
    path=tmp_path/'cloud.xlsx';wb.save(path)
    cells,warnings=from_workbook(path)
    assert cells[0]['text']=='0012'
    assert cells[1]['text']==''
    assert any('无显示缓存' in w for w in warnings)


def test_api_review_save_export_and_revision(client,tmp_path):
    r=client.post('/api/jobs',files={'file':('photo.png',photo(),'image/png')},data={'geometry':'false'})
    assert r.status_code==200
    job=r.json();jid=job['id']
    assert (tmp_path/jid/'source.upload').read_bytes()==photo()
    assert client.get(f'/api/jobs/{jid}/excel').status_code==409
    c=cell(text='=2+2',confirmed=True)
    r=client.put(f'/api/jobs/{jid}',json={'revision':0,'cells':[c]})
    assert r.status_code==200
    assert client.put(f'/api/jobs/{jid}',json={'revision':0,'cells':[c]}).status_code==409
    assert client.get(f'/api/jobs/{jid}/excel').status_code==200
    wb=load_workbook(io.BytesIO(client.get(f'/api/jobs/{jid}/excel').content))
    assert wb.active['A1'].value=='=2+2' and wb.active['A1'].data_type=='s'
    assert client.post(f'/api/jobs/{jid}/warp',json={'points':[[0,0],[300,0],[300,190],[0,190]]}).status_code==409
    assert client.post(f'/api/jobs/{jid}/recognize',json={}).status_code==409


def test_failed_ocr_is_not_empty_success(client,monkeypatch):
    def fail(*args,**kwargs):
        raise ValueError('无登录凭证')
    monkeypatch.setattr(module.providers,'recognize',fail)
    job=client.post('/api/jobs',files={'file':('x.png',photo(),'image/png')}).json()
    jid=job['id']
    assert client.post(f'/api/jobs/{jid}/recognize',json={}).status_code==200
    for _ in range(100):
        state=client.get(f'/api/jobs/{jid}').json()
        if state['status']!='recognizing':break
        time.sleep(.01)
    assert state['status']=='failed' and state['error']=='表格识别失败：无登录凭证'
    assert state['failed_stage']=='table'
    assert state['cells']==[]


def test_host_and_foreign_origin_blocked(client):
    assert client.get('/api/health',headers={'Host':'evil.example'}).status_code==403
    assert client.post('/api/jobs',headers={'Origin':'https://evil.example'}).status_code==403


def test_auto_geometry_preserves_entire_canvas(tmp_path):
    from PIL import ImageDraw
    import cv2
    import numpy as np
    import imaging
    image=Image.new('RGB',(1000,800),'#444444')
    draw=ImageDraw.Draw(image)
    draw.polygon([(150,100),(850,70),(900,700),(100,700)],fill='white',outline='black',width=4)
    draw.text((10,20),'OUTSIDE NOTE',fill='white')
    source=tmp_path/'source.png';target=tmp_path/'processed.png';image.save(source)
    matrix,warnings=imaging.preprocess(source,target,True,'natural')
    assert any('完整画面' in w for w in warnings)
    with Image.open(target) as result:
        corners=cv2.perspectiveTransform(np.float32([[[0,0],[999,0],[999,799],[0,799]]]),np.array(matrix))[0]
        assert (corners>=0).all()
        assert (corners[:,0]<result.width).all() and (corners[:,1]<result.height).all()


def test_login_starts_official_cli_once(monkeypatch):
    from types import SimpleNamespace
    import providers
    calls=[]
    monkeypatch.setattr(providers,'LOGIN_PROCESS',None)
    def launch(*args,**kwargs):
        calls.append((args,kwargs))
        return SimpleNamespace(poll=lambda:None,stdout=io.StringIO(""),wait=lambda:0)
    monkeypatch.setattr(providers.subprocess,'Popen',launch)
    providers.start_login()
    providers.start_login()
    assert len(calls)==1
    assert calls[0][0][0]==[str(providers.CLI),'auth','login','--no-browser']
    assert calls[0][1]['env']['LOCALAPPDATA']==str(providers.ROOT/'data/auth')


def test_auth_api_returns_safe_state_and_launches_login(client,monkeypatch):
    import providers
    monkeypatch.setattr(providers,'cli',lambda *a,**k:'{"authenticated":true,"token":"never-return-this"}')
    monkeypatch.setattr(providers,'login_running',lambda:False)
    started=[]
    monkeypatch.setattr(providers,'start_login',lambda:started.append(True))
    state=client.get('/api/auth/status').json()
    assert state['logged_in'] and 'token' not in state and 'never-return-this' not in str(state)
    assert client.post('/api/auth/login').status_code==200
    assert started==[True]


def test_line_coordinates_require_matching_cloud_structure(tmp_path):
    from PIL import ImageDraw
    import imaging
    image=Image.new('RGB',(600,400),'white');draw=ImageDraw.Draw(image)
    for x in [40,210,380,560]:draw.line((x,40,x,350),fill='black',width=3)
    for y in [40,150,240,350]:draw.line((40,y,560,y),fill='black',width=3)
    path=tmp_path/'table.png';image.save(path)
    cells=[cell(id=f'{r}-{c}',row=r,col=c) for r in range(1,4) for c in range(1,4)]
    located,warnings=imaging.locate_cloud_cells(path,cells)
    assert all(len(c['polygon'])==8 for c in located)
    assert any('候选' in w for w in warnings)
    assert all(c['text']=='0012' and not c['confirmed'] for c in located)
    mismatched=[dict(c) for c in cells]+[cell(id='extra',row=4)]
    fallback,warnings=imaging.locate_cloud_cells(path,mismatched)
    assert all(not c['polygon'] for c in fallback)
    assert any('不一致' in w for w in warnings)


def test_export_allows_partial_confirmation_without_changing_review_status(client):
    j=client.post('/api/jobs',files={'file':('draft.png',photo(),'image/png')}).json()
    r=client.put(f'/api/jobs/{j["id"]}',json={'revision':0,'cells':[cell(text='0012',confirmed=False)]})
    assert r.status_code==200
    assert client.get(f'/api/jobs/{j["id"]}/excel').status_code==200
    assert client.get(f'/api/jobs/{j["id"]}').json()['cells'][0]['confirmed'] is False
    draft=client.get(f'/api/jobs/{j["id"]}/excel?draft=true')
    assert draft.status_code==200
    wb=load_workbook(io.BytesIO(draft.content))
    assert wb.active['A1'].value=='0012' and wb.active['A1'].data_type=='s'
    assert wb.sheetnames==['识别表格','原图']


def test_outside_signature_requires_independent_position(tmp_path):
    from PIL import ImageDraw
    import imaging
    image=Image.new('RGB',(600,400),'white');draw=ImageDraw.Draw(image)
    for x in [40,210,380,560]:draw.line((x,40,x,300),fill='black',width=3)
    for y in [40,130,220,300]:draw.line((40,y,560,y),fill='black',width=3)
    p=tmp_path/'table.png';image.save(p)
    cells=[cell(id=f'{r}-{c}',row=r,col=c) for r in range(1,4) for c in range(1,4)]
    cells.append(cell(id='signature',row=4,col=1,colspan=3,text='签收人签字：'))
    line=dict(text='签收人签字：',polygon=[400,325,550,325,550,350,400,350])
    located,_=imaging.locate_cloud_cells(p,cells,[line])
    assert all(len(c['polygon'])==8 for c in located)
    assert located[-1]['kind']=='outside' and located[-1]['polygon']==line['polygon']
    rejected,_=imaging.locate_cloud_cells(p,cells,[dict(line,polygon=[400,100,550,100,550,120,400,120])])
    assert all(not c['polygon'] for c in rejected)


def test_auxiliary_scores_do_not_replace_text_and_survive_save(client):
    import assessment
    cells=[cell(text='0012',score=None),cell(id='empty',col=2,text='',score=None)]
    assessed=assessment.attach_candidates(cells,{'0':dict(text='0017',score=.731)})
    assert assessed[0]['text']=='0012' and assessed[0]['score'] is None
    assert assessed[0]['candidate_score']==.731 and assessed[0]['candidate_text']=='0017'
    assert assessed[1]['candidate_score'] is None
    j=client.post('/api/jobs',files={'file':('x.png',photo(),'image/png')}).json()
    result=client.put(f'/api/jobs/{j["id"]}',json={'revision':0,'cells':assessed})
    assert result.status_code==200
    assert result.json()['cells'][0]['candidate_score']==.731
    draft=client.get(f'/api/jobs/{j["id"]}/excel?draft=true')
    wb=load_workbook(io.BytesIO(draft.content))
    assert wb.active['A1'].value=='0012' and wb.active['B1'].value is None


def test_original_registration_rejects_featureless_images(tmp_path):
    import imaging
    p=tmp_path/'white.png';Image.new('RGB',(300,300),'white').save(p)
    matrix,evidence=imaging.register_original(p,p)
    assert matrix is None and evidence['reason']


def test_successful_enhancement_is_reused_after_conversion_failure(client,monkeypatch):
    import shutil
    calls=[]
    def enhance(args,*a,**k):
        calls.append('enhance');shutil.copyfile(args[2],args[-1]);return ''
    def convert(image,*a,**k):
        calls.append('convert')
        if calls.count('convert')==1:raise ValueError('官方服务限流')
        target=image.parent/'cloud.xlsx'
        wb=Workbook();wb.active['A1']='0012';wb.save(target)
        return from_workbook(target),None
    monkeypatch.setattr(module.providers,'cli',enhance)
    monkeypatch.setattr(module.providers,'recognize',convert)
    monkeypatch.setattr(module.assessment,'assess',lambda path,cells:(cells,[],None,{'registration':{}}))
    j=client.post('/api/jobs',files={'file':('x.png',photo(),'image/png')}).json()
    def attempt():
        client.post(f'/api/jobs/{j["id"]}/recognize',json={'cloud_enhance':True})
        for _ in range(150):
            state=client.get(f'/api/jobs/{j["id"]}').json()
            if state['status']!='recognizing':return state
            time.sleep(.01)
        pytest.fail('worker did not complete')
    failed=attempt()
    assert failed['failed_stage']=='table' and failed['status']=='failed'
    recovered=attempt()
    assert recovered['status']=='review' and recovered['cells'][0]['text']=='0012'
    assert calls==['enhance','convert','convert']


def test_cli_error_is_classified_without_exposing_credentials(monkeypatch):
    import providers
    from types import SimpleNamespace
    calls=[]
    def invoke(*a,**k):
        calls.append(1)
        return SimpleNamespace(returncode=1,stdout='',stderr='HTTP 429 rate limit Authorization: Bearer private-test-token')
    monkeypatch.setattr(providers.subprocess,'run',invoke)
    with pytest.raises(ValueError) as exc:providers.cli(['image','enhance','test.png'])
    assert '限流' in str(exc.value) and 'private-test-token' not in str(exc.value)
    assert len(calls)==1
