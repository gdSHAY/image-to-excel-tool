from pathlib import Path
import math
import re
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter


def validate(cells):
    if len(cells) > 30000:
        raise ValueError('超过 30000 格限制')
    occupied = set()
    ids = set()
    for c in cells:
        if c['id'] in ids:
            raise ValueError('重复单元格 ID')
        ids.add(c['id'])
        r, col, rs, cs = (int(c[k]) for k in ('row', 'col', 'rowspan', 'colspan'))
        if min(r, col, rs, cs) < 1 or r+rs > 10001 or col+cs > 501 or rs*cs > 30000:
            raise ValueError('行列或合并范围无效（最多10000行、500列）')
        if len(c['text']) > 32767:
            raise ValueError('单元格文字超过 Excel 上限')
        if any(ord(ch) < 32 and ch not in '\n\r\t' for ch in c['text']):
            raise ValueError('文字包含 Excel 不支持的控制字符')
        score = c.get('score')
        if score is not None and (not math.isfinite(score) or not 0 <= score <= 1):
            raise ValueError('分数必须在0到1之间')
        for rr in range(r, r+rs):
            for cc in range(col, col+cs):
                if (rr,cc) in occupied:
                    raise ValueError('单元格/合并范围重叠，请修改结构后保存')
                occupied.add((rr,cc))


def from_textin(payload):
    if payload.get('code') != 200:
        raise ValueError(f"TextIn 错误 {payload.get('code')}：{payload.get('message', '未知原因')}")
    result = payload.get('result', {})
    areas = result.get('tables', [])
    cells, warnings = [], []
    # Preserve area order by geometry, including titles and notes.
    areas = sorted(areas, key=lambda a: (min(a.get('position', [0])[1::2] or [0]), a.get('area_index', 0)))
    offset = 0
    for area in areas:
        table_cells = area.get('table_cells', [])
        if table_cells:
            for raw in table_cells:
                # Official sample uses one-based inclusive start/end indexes.
                r, col = int(raw['start_row']), int(raw['start_col'])
                scores = [float(l['score']) for l in raw.get('lines', []) if l.get('score') is not None]
                cells.append(dict(id=str(len(cells)), row=offset+r, col=col,
                    rowspan=int(raw['end_row'])-r+1, colspan=int(raw['end_col'])-col+1,
                    text=str(raw.get('text', '')), score=min(scores) if scores else None,
                    polygon=raw.get('position', []), confirmed=False, kind='cell'))
            offset = max(c['row']+c['rowspan']-1 for c in cells)
        else:
            for line in sorted(area.get('lines', []), key=lambda l: min(l.get('position', [0])[1::2] or [0])):
                offset += 1
                cells.append(dict(id=str(len(cells)), row=offset, col=1, rowspan=1, colspan=1,
                    text=str(line.get('text','')), score=line.get('score'), polygon=line.get('position', []),
                    confirmed=False, kind='outside'))
    if not cells:
        raise ValueError('API 没有返回可识别的单元格或文字，不能当作空表成功')
    if any(c['kind'] == 'outside' for c in cells):
        warnings.append('表格外文字按区域顺序保留；请核对标题、备注与相邻表格的位置。')
    validate(cells)
    return cells, warnings


def from_workbook(path):
    wb = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    cells, warnings = [], ['扫描全能王转换未返回格子坐标或分数。请逐格看图校对，必要时手动绑定区域；不能保证前导零已被云端保留。']
    offset = 0
    for s, sheet in enumerate(wb):
        if sheet.max_row * sheet.max_column > 30000:
            raise ValueError('云端工作簿过大，请分表识别')
        anchors, covered = {}, set()
        for merged in sheet.merged_cells.ranges:
            anchors[(merged.min_row,merged.min_col)] = (merged.max_row-merged.min_row+1,merged.max_col-merged.min_col+1)
            covered.update((r,c) for r in range(merged.min_row,merged.max_row+1) for c in range(merged.min_col,merged.max_col+1) if (r,c)!=(merged.min_row,merged.min_col))
        for row in sheet:
            for raw in row:
                if (raw.row,raw.column) in covered:
                    continue
                value = raw.value
                if raw.data_type == 'f':
                    value = cached.worksheets[s][raw.coordinate].value
                    if value is None:
                        warnings.append(f'{sheet.title}!{raw.coordinate} 公式无显示缓存，留空待校对。')
                text = '' if value is None else str(value)
                if isinstance(value, (int,float)) and raw.number_format and set(raw.number_format) == {'0'}:
                    text = str(int(value)).zfill(len(raw.number_format))
                rs,cs = anchors.get((raw.row,raw.column),(1,1))
                cells.append(dict(id=str(len(cells)),row=offset+raw.row,col=raw.column,rowspan=rs,colspan=cs,
                    text=text,score=None,polygon=[],confirmed=False,kind='cell'))
        offset += sheet.max_row
    if not any(c['text'] for c in cells):
        raise ValueError('转换结果没有文字，需重新识别或检查原图')
    validate(cells)
    return cells,warnings


def excel_value(text):
    """Convert plain numbers only; keep identifiers and expressions literal."""
    value = text.strip()
    if not re.fullmatch(r'[+-]?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?', value):
        return text, '@'
    # Excel stores at most 15 significant digits. Never round long identifiers.
    if len(Decimal(value).as_tuple().digits) > 15:
        return text, '@'
    if '.' in value:
        number = float(value)
        if not math.isfinite(number) or (number == 0 and Decimal(value) != 0):
            return text, '@'
        return number, '0.' + '0' * len(value.split('.')[1])
    return int(value), '0'


def export(cells, original: Path, target: Path):
    validate(cells)
    wb = Workbook()
    ws = wb.active
    ws.title = '识别表格'
    edge = Side(style='thin', color='888888')
    widths, heights = {}, {}
    for c in cells:
        item = ws.cell(c['row'],c['col'])
        item.value, item.number_format = excel_value(c['text'])
        if isinstance(item.value, str):
            item.data_type = 's'  # Formula-like text and leading zeros stay literal.
        item.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
        item.border = Border(left=edge,right=edge,top=edge,bottom=edge)
        rs,cs = c['rowspan'],c['colspan']
        if rs > 1 or cs > 1:
            ws.merge_cells(start_row=c['row'],start_column=c['col'],end_row=c['row']+rs-1,end_column=c['col']+cs-1)
        lines = c['text'].split('\n')
        width = min(48,max(10,max((sum(2 if ord(ch)>255 else 1 for ch in line) for line in lines),default=0)/cs+2))
        for col in range(c['col'],c['col']+cs):
            widths[col] = max(widths.get(col,10),width)
        heights[c['row']] = max(heights.get(c['row'],24), 18*max(len(lines), math.ceil(len(c['text']) / max(1,width*cs-2))))
    for col,width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for row,height in heights.items():
        ws.row_dimensions[row].height = min(409,height+6)
    source = wb.create_sheet('原图')
    source.add_image(XLImage(str(original)), 'A1')
    wb.save(target)
